"""Build lds3_survey_clean.csv and the LDS 3.0 Excel export from the newest Qualtrics export.

Run from anywhere:  python scripts/lds-survey/build_lds3_data.py
LDS 3.0 is a separate program; its data never mixes with lds_survey_clean.csv.
"""
import csv
import glob
import os
import re
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

SRC_DIR = r"C:\Users\USER\Desktop\NDR Files\LDS survey\LDS 3.0"
REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT_DIR = os.path.join(REPO, "static", "lds-survey-dashboard")
OUT_CSV = os.path.join(OUT_DIR, "lds3_survey_clean.csv")
OUT_XLSX = os.path.join(OUT_DIR, "LDS3_Survey_Export.xlsx")

DEPT_NORMALIZE = {"ND Research": "NDR"}
# Names that must never reach the page live in a git-ignored file so they are never published:
# one per line, "Full Name => replacement" (replacement defaults to "the facilitator"). Longest names first.
NAMES_FILE = os.path.join(os.path.dirname(__file__), "private_names.txt")


def load_redactions():
    if not os.path.exists(NAMES_FILE):
        print(f"WARNING: {NAMES_FILE} not found; comments will not be name-redacted")
        return []
    out = []
    for line in open(NAMES_FILE, encoding="utf-8"):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        name, _, repl = (part.strip() for part in line.partition("=>"))
        out.append((re.compile(rf"\b{re.escape(name)}\b"), repl or "the facilitator"))
    return out


REDACT = load_redactions()

FIELDS = ["StartDate", "Q1", "Q2", "Q3_1", "Q3_2", "Q3_3", "Q4_1", "Q5_1", "Q17",
          "Q4_1_numeric", "Q5_1_numeric", "Q4_1_category", "Q5_1_category"]


def newest_export():
    files = [f for f in glob.glob(os.path.join(SRC_DIR, "*.xlsx")) if not os.path.basename(f).startswith("~$")]
    if not files:
        raise SystemExit(f"No .xlsx export found in {SRC_DIR}")
    return max(files, key=os.path.getmtime)


def to_num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def category(score):
    if score is None:
        return "No Response"
    return "Detractor" if score <= 6 else "Passive" if score <= 8 else "Promoter"


def redact(text):
    for pattern, repl in REDACT:
        text = pattern.sub(repl, text)
    if text[:1].islower() and text.startswith("the facilitator"):
        text = "T" + text[1:]
    return text.strip()


def fmt_score(n):
    if n is None:
        return ""
    return str(int(n)) if n == int(n) else str(n)


def load_rows(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    raw = list(ws.iter_rows(values_only=True))
    idx = {name: i for i, name in enumerate(raw[0])}
    rows, preview = [], 0
    for r in raw[2:]:  # row 0 = codes, row 1 = question text
        if r[idx["Status"]] == "Survey Preview":
            preview += 1
            continue
        q4, q5 = to_num(r[idx["Q4_1"]]), to_num(r[idx["Q5_1"]])
        start = r[idx["StartDate"]]
        rows.append({
            "StartDate": start.strftime("%Y-%m-%d %H:%M:%S") if isinstance(start, datetime) else (start or ""),
            "Q1": DEPT_NORMALIZE.get(r[idx["Q1"]] or "", r[idx["Q1"]] or ""),
            "Q2": r[idx["Q2"]] or "",
            "Q3_1": r[idx["Q3_1"]] or "",
            "Q3_2": r[idx["Q3_2"]] or "",
            "Q3_3": r[idx["Q3_3"]] or "",
            "Q4_1": fmt_score(q4),
            "Q5_1": fmt_score(q5),
            "Q17": redact(r[idx["Q17"]] or ""),
            "Q4_1_numeric": "" if q4 is None else q4,
            "Q5_1_numeric": "" if q5 is None else q5,
            "Q4_1_category": category(q4),
            "Q5_1_category": category(q5),
        })
    return rows, preview


def write_csv(rows):
    with open(OUT_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows):
    # Header names match what the page's "Upload Updated Data" parser looks for.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Survey Data (Editable)"
    ws["A1"] = "INSTRUCTIONS: You can edit the data below. When done, save this file and upload it back to the dashboard."
    ws["A1"].font = Font(bold=True, size=12, color="C00000")
    ws.merge_cells("A1:I1")
    ws["A2"] = f"LDS 3.0 | Last Updated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(italic=True, size=10)
    ws.merge_cells("A2:I2")
    headers = ["Start Date", "Department", "Training Session", "Content Relevance", "Duration Appropriate",
               "Future Usability", "Training Score (0-10)", "Facilitator Score (0-10)", "Comments"]
    fill = PatternFill(start_color="0C2340", end_color="0C2340", fill_type="solid")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.fill, cell.font = fill, Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r_i, row in enumerate(rows, 5):
        values = [row["StartDate"], row["Q1"], row["Q2"], row["Q3_1"], row["Q3_2"], row["Q3_3"],
                  row["Q4_1_numeric"] if row["Q4_1_numeric"] != "" else None,
                  row["Q5_1_numeric"] if row["Q5_1_numeric"] != "" else None,
                  row["Q17"] or None]
        for c, v in enumerate(values, 1):
            ws.cell(row=r_i, column=c, value=v)
    for col, width in zip("ABCDEFGHI", [20, 30, 55, 18, 20, 18, 14, 14, 70]):
        ws.column_dimensions[col].width = width
    wb.save(OUT_XLSX)


if __name__ == "__main__":
    src = newest_export()
    rows, preview = load_rows(src)
    write_csv(rows)
    write_xlsx(rows)
    print(f"Source: {os.path.basename(src)}")
    print(f"Dropped {preview} preview row(s); wrote {len(rows)} responses")
    print(f"  -> {OUT_CSV}")
    print(f"  -> {OUT_XLSX}")
